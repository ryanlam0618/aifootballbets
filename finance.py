import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from config import settings

def calculate_kelly_stake(prob: float, odds: float, bankroll: float) -> dict:
    if prob <= 0 or odds <= 1:
        return {"stake": 0, "pct": 0, "ev": 0}

    b = odds - 1
    q = 1 - prob
    f = (b * prob - q) / b
    f_adj = f * settings.KELLY_FRACTION
    ev = (prob * b) - (1 - prob)
    
    if f_adj <= 0 or ev < settings.MIN_EDGE:
        return {"stake": 0, "pct": 0, "ev": ev}
    
    stake = bankroll * f_adj
    return {"stake": stake, "pct": f_adj, "ev": ev}

class ExcelLogger:
    def __init__(self):
        self.filepath = settings.EXCEL_FILEPATH
        self.columns = [
            "Date", "League", "Home", "Away", "Market", "Selection", 
            "Odds", "Model Prob", "EV", "Kelly %", "Stake ($)", 
            "Result (Win/Loss)", "Profit"
        ]
        self._init_file()

    def _init_file(self):
        # 嘗試建立目錄與檔案
        if not os.path.exists(self.filepath):
            try:
                directory = os.path.dirname(self.filepath)
                if directory and not os.path.exists(directory):
                    os.makedirs(directory)
                    
                df = pd.DataFrame(columns=self.columns)
                df.to_excel(self.filepath, index=False)
                print(f"📄 已建立新記錄檔: {self.filepath}")
            except Exception as e:
                print(f"❌ 建立檔案失敗: {e}")
        else:
            print(f"✅ 讀取現有記錄檔: {self.filepath}")

    def log_bet(self, match_info, bet_info, stake_info):
        new_row = {
            "Date": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "League": match_info.get("league"),
            "Home": match_info.get("home"),
            "Away": match_info.get("away"),
            "Market": bet_info.get("market"),
            "Selection": bet_info.get("selection"),
            "Odds": bet_info.get("odds"),
            "Model Prob": round(bet_info.get("model_probability"), 3),
            "EV": round(stake_info["ev"], 3),
            "Kelly %": f"{stake_info['pct']*100:.2f}%",
            "Stake ($)": round(stake_info["stake"], 2),
            "Result (Win/Loss)": "", 
            "Profit": ""
        }
        
        try:
            df_new = pd.DataFrame([new_row])
            
            if os.path.exists(self.filepath):
                # 修正部分：先用 openpyxl 讀取一次來計算開始行數 (start_row)
                wb = load_workbook(self.filepath)
                # 取得目前工作表的最大行數
                start_row = wb.active.max_row
                
                # 寫入模式：使用 overlay 模式，從 start_row 開始寫，不寫入 header
                with pd.ExcelWriter(self.filepath, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                    # 這裡不需要再設定 writer.book = ...
                    df_new.to_excel(writer, index=False, header=False, startrow=start_row)
            else:
                # 如果檔案不存在，直接寫入 (包含 header)
                df_new.to_excel(self.filepath, index=False)
                
            print(f"💾 記錄已成功儲存至 Excel")
            
        except PermissionError:
            print("❌ 寫入失敗: 請先關閉 Excel 檔案後再試！")
        except Exception as e:
            print(f"❌ 寫入失敗: {e}")

    def show_stats(self):
        if not os.path.exists(self.filepath):
            return
        try:
            df = pd.read_excel(self.filepath)
            # 確保欄位存在再過濾
            if "Result (Win/Loss)" in df.columns:
                settled = df[df["Result (Win/Loss)"].notna() & (df["Result (Win/Loss)"] != "")]
                if not settled.empty:
                    wins = settled[settled["Result (Win/Loss)"].astype(str).str.upper().str.contains("W")]
                    win_rate = len(wins) / len(settled)
                    print(f"📊 目前勝率: {win_rate*100:.1f}% ({len(wins)}/{len(settled)})")
        except Exception:
            pass